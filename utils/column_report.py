"""Informe del análisis momento-curvatura (COLUMN) en Word y PDF.

El contenido se arma una sola vez como un documento neutro (`ReportDocument`: títulos, párrafos,
viñetas, tablas e imágenes) y dos escritores lo vuelcan a .docx (python-docx) y a .pdf (fpdf2).
Así ambos formatos tienen exactamente el mismo contenido.

Dependencias opcionales: python-docx y fpdf2 (ver requirements.txt). Se importan al escribir, de
modo que la app funciona aunque falten; el error indica qué instalar.
"""

import datetime
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from calculations.structural.column_analysis import MomentCurvatureResult
from visualization.column_report_plots import (
    PHI_UNIT,
    PngImage,
    interaction_png,
    materials_png,
    moment_curvature_png,
    section_png,
    variants_png,
)

# =========================================================================================
# Documento neutro
# =========================================================================================


@dataclass
class Heading:
    text: str
    level: int = 1


@dataclass
class Paragraph:
    text: str


@dataclass
class Bullets:
    items: list


@dataclass
class Table:
    headers: list
    rows: list
    widths: list | None = None  # anchos relativos de las columnas


@dataclass
class Picture:
    image: PngImage
    caption: str = ""
    width_cm: float = 15.0


@dataclass
class ReportDocument:
    title: str
    subtitle: str = ""
    blocks: list = field(default_factory=list)

    def add(self, *blocks):
        self.blocks.extend(blocks)


def plain(text: str) -> str:
    """Quita el énfasis de Markdown (**negrita**) para los formatos que no lo interpretan."""
    return text.replace("**", "")


# =========================================================================================
# Interpretación (la usan la página y el informe)
# =========================================================================================


def interpretation_lines(result: MomentCurvatureResult) -> list:
    """Frases de interpretación de los resultados (Markdown, una por elemento)."""
    records = {record.key: record for record in result.limit_states}
    maximum, final = result.maximum, result.states[-1]
    ultimate_duct = result.ductility_of(final.phi)
    lines = []
    ratio = maximum.moment / result.ideal_moment
    if ratio >= 1.0:
        lines.append(
            f"La sección alcanza **Mmax = {maximum.moment / 1e6:,.1f} kN·m**, que es {ratio:.2f} veces la resistencia ideal Mi: el exceso "
            "proviene sobre todo del endurecimiento del acero y del confinamiento del núcleo, que el método de la norma no considera."
        )
    else:
        lines.append(
            f"La sección alcanza **Mmax = {maximum.moment / 1e6:,.1f} kN·m**, que es solo {ratio:.2f} veces la resistencia ideal Mi: "
            "revise los datos, porque normalmente Mmax supera a Mi."
        )
    first_yield = records["first_yield"]
    if first_yield.reached and first_yield.curvature <= maximum.phi:
        lines.append(
            f"El acero traccionado fluye (φ'y = {result.yield_curvature_first * PHI_UNIT:.2f} µrad/mm) antes del momento máximo, "
            "de modo que la curva tiene un tramo de fluencia. Con φy = φ'y·Mi/My (Ec. 7.60) la ductilidad se mide como μφ = φ/φy."
        )
    elif first_yield.reached:
        lines.append("El momento máximo ocurre antes de la primera fluencia del acero traccionado: la respuesta está controlada por la compresión del concreto.")
    else:
        lines.append("El acero traccionado **no fluye** antes del estado límite que termina el análisis: la sección está controlada por compresión y no se define ductilidad de curvatura.")
    spalling = records.get("cover_spalling")
    if spalling is not None and spalling.reached:
        drop = 1.0 - spalling.moment / maximum.moment
        lines.append(
            f"El recubrimiento se descascara a φ = {spalling.curvature * PHI_UNIT:.1f} µrad/mm (εspall = {result.model.concrete.spall_strain:g}); "
            f"en ese punto el momento está {drop:.0%} por debajo del máximo. Después la capacidad depende del núcleo confinado."
        )
    lines.append(
        f"El análisis terminó por **{result.termination_message.lower()}** a φ = {final.phi * PHI_UNIT:.1f} µrad/mm"
        + (f" (μφ = {ultimate_duct:.1f})" if ultimate_duct is not None else "")
        + f", con M = {final.moment / 1e6:,.1f} kN·m ({final.moment / maximum.moment:.0%} de Mmax). "
        "Si cambia el criterio de terminación, la curvatura última cambia."
    )
    if result.termination_key == "hoop_fracture":
        lines.append("La fractura del estribo se estima con un balance de energía (Mander, §3.8.2) cuyas constantes (105 MJ/m³ para el acero y 0.017√f'c para el concreto) son empíricas.")
    if result.termination_key == "steel_ultimate":
        lines.append("Rige la deformación última del acero εsu; en una columna real el confinamiento o el pandeo de las barras suelen limitar antes la curvatura útil.")
    lines.append(
        "**Limitaciones:** respuesta monotónica sobre la curva esqueleto (sin descarga de fibras), cortante y deslizamiento ignorados, "
        "reconstrucción de COLUMN y no su traducción literal."
    )
    return lines


# =========================================================================================
# Contenido del informe
# =========================================================================================


def _num(value, decimals=2, thousands=True):
    if value is None:
        return "—"
    return f"{value:,.{decimals}f}" if thousands else f"{value:.{decimals}f}"


def build_column_report(result: MomentCurvatureResult, arrangement=None, interaction=None, variants=None, title="", author="") -> ReportDocument:
    """Arma el informe.

    arrangement: dict con 'octagonal', 'legs_x', 'legs_y' y 'description' (cómo se dibuja y describe el estribo).
    interaction: dict con 'curve' (CodeInteractionCurve), 'points' (lista de InteractionPoint) y 'current_axial_kn'.
    variants:    lista de VariantResult de la comparación de configuraciones.
    """
    arrangement = arrangement or {}
    model, settings = result.model, result.settings
    section, concrete, confinement, steel = model.section, model.concrete, model.confinement, model.steel
    transverse = section.transverse
    maximum, final = result.maximum, result.states[-1]
    gross_strength = concrete.unconfined_strength * section.gross_area_mm2

    doc = ReportDocument(
        title=title or "Análisis momento–curvatura de una sección de concreto armado",
        subtitle="Reconstrucción monotónica del análisis COLUMN (Mander, 1984)"
        + (f" — {author}" if author else "")
        + f" — {datetime.date.today():%d/%m/%Y}",
    )

    # ---------- 1. Datos ----------
    doc.add(Heading("1. Datos de entrada"))
    doc.add(Heading("1.1 Geometría y refuerzo", 2))
    doc.add(Table(["Parámetro", "Valor"], [
        ["Ancho b (mm)", _num(section.width_mm, 1)],
        ["Peralte h (mm), flexión respecto al eje x", _num(section.depth_mm, 1)],
        ["Recubrimiento libre hasta el estribo (mm)", _num(section.clear_cover_mm, 1)],
        ["Núcleo b″ × d″ (mm), al eje del estribo", f"{section.core_width_mm:.1f} × {section.core_depth_mm:.1f}"],
        ["Área bruta Ag (mm²)", _num(section.gross_area_mm2, 0)],
        ["Acero longitudinal As (mm²) / número de barras", f"{section.steel_area_mm2:,.0f} / {section.number_of_bars}"],
        ["Cuantía As/Ag", f"{section.steel_area_mm2 / section.gross_area_mm2:.2%}"],
        ["Cuantía respecto al núcleo ρcc", f"{section.longitudinal_ratio_core:.4f}"],
    ], widths=[3, 1.4]))
    doc.add(Table(["Capa (de la cara comprimida a la inferior)", "Profundidad (mm)", "N° de barras", "Diámetro (mm)", "Área (mm²)"],
                  [[str(i), _num(l.depth_mm, 1), str(l.count), _num(l.diameter_mm, 1), _num(l.area_mm2, 0)]
                   for i, l in enumerate(section.layers, start=1)], widths=[3, 1.5, 1.3, 1.3, 1.3]))
    doc.add(Picture(section_png(section, arrangement.get("octagonal", False), arrangement.get("legs_x", 2), arrangement.get("legs_y", 2)),
                    "Figura 1. Corte de la sección con el refuerzo ingresado (a escala; cotas en mm).", width_cm=12.5))

    doc.add(Heading("1.2 Estribos y confinamiento", 2))
    doc.add(Table(["Parámetro", "Valor"], [
        ["Diámetro del estribo (mm)", _num(transverse.bar_diameter_mm, 1)],
        ["Espaciamiento s (mm)", _num(transverse.spacing_mm, 1)],
        ["fyh (MPa)", _num(transverse.yield_stress, 0)],
        ["Configuración", arrangement.get("description", "—")],
        ["Acero transversal paralelo a x / a y (mm²)", f"{transverse.area_x_mm2:,.1f} / {transverse.area_y_mm2:,.1f}"],
        ["Coeficiente de efectividad ke", f"{confinement.effectiveness:.3f}"],
        ["ρx / ρy", f"{confinement.rho_x:.4f} / {confinement.rho_y:.4f}"],
        ["Confinamiento lateral f'lx / f'ly (MPa)", f"{confinement.lateral_stress_x:.2f} / {confinement.lateral_stress_y:.2f}"],
        ["Resistencia confinada f'cc (MPa) y razón f'cc/f'co", f"{confinement.confined_strength:.2f}  (×{confinement.strength_ratio:.3f})"],
        ["Deformación en f'cc, εcc", f"{confinement.confined_strain:.5f}"],
    ], widths=[3, 1.4]))

    doc.add(Heading("1.3 Materiales y parámetros del análisis", 2))
    doc.add(Table(["Parámetro", "Valor"], [
        ["Concreto: f'co (MPa) / εco", f"{concrete.unconfined_strength:.2f} / {concrete.strain_at_peak:.5f}"],
        ["Concreto: εspall / f't (MPa) / Ec (MPa)", f"{concrete.spall_strain:g} / {concrete.resolved_tensile_strength:.2f} / {concrete.resolved_elastic_modulus:,.0f}"],
        ["Acero: fy / fsu (MPa)", f"{steel.yield_stress:.0f} / {steel.ultimate_stress:.0f}"],
        ["Acero: Es / Esh (MPa)", f"{steel.elastic_modulus:,.0f} / {steel.hardening_modulus:,.0f}"],
        ["Acero: εsh / εsu", f"{steel.hardening_strain:g} / {steel.ultimate_strain:g}"],
        ["Carga axial P (kN), compresión positiva", f"{settings.axial_load_n / 1e3:,.1f}  (P/f'co·Ag = {settings.axial_load_n / gross_strength:.3f})"],
        ["Láminas de concreto", str(len(model.mesh.y))],
        ["Incrementos de curvatura", "COLUMN (0.2, 0.5 y 1.0 φy)" if settings.increment_schedule == "column" else f"uniformes, {settings.uniform_increment * PHI_UNIT:g} µrad/mm"],
        ["Terminación", "; ".join(filter(None, [
            "fractura del estribo (balance de energía)" if settings.use_hoop_fracture else "",
            f"εcu del núcleo = {settings.core_ultimate_strain:g}" if settings.core_ultimate_strain else "",
            f"φ/φy ≤ {settings.max_curvature_ductility:g}" if settings.max_curvature_ductility else "",
            f"M ≥ {settings.moment_drop_ratio:.0%} de Mmax" if settings.moment_drop_ratio else "",
            "εsu del acero"]))],
        ["Tolerancia de equilibrio axial", f"{settings.equilibrium_tolerance:g} de f'co·Ag"],
    ], widths=[3, 1.4]))

    # ---------- 2. Resultados ----------
    doc.add(Heading("2. Resultados"))
    ultimate_duct = result.ductility_of(final.phi)
    doc.add(Table(["Magnitud", "Valor"], [
        ["Momento máximo Mmax (kN·m)", _num(maximum.moment / 1e6, 1)],
        ["Curvatura en Mmax (µrad/mm)", _num(maximum.phi * PHI_UNIT, 1)],
        ["Resistencia ideal Mi (kN·m), método de la norma", _num(result.ideal_moment / 1e6, 1)],
        ["Primera fluencia: φ'y (µrad/mm) / My (kN·m)", f"{_num(None if result.yield_curvature_first is None else result.yield_curvature_first * PHI_UNIT)} / {_num(None if result.yield_moment_first is None else result.yield_moment_first / 1e6, 1)}"],
        ["Curvatura de fluencia φy = φ'y·Mi/My (µrad/mm)", _num(None if result.yield_curvature is None else result.yield_curvature * PHI_UNIT)],
        [f"Curvatura al terminar φu (µrad/mm) — {result.termination_message}", _num(final.phi * PHI_UNIT, 1)],
        ["Ductilidad de curvatura μφ = φu/φy", _num(ultimate_duct, 1)],
        ["Incrementos / iteraciones de Newton máximas", f"{len(result.states)} / {max(s.iterations for s in result.states)}"],
    ], widths=[3, 1.4]))
    doc.add(Paragraph("La curvatura última y la ductilidad dependen del estado límite que termina el análisis; no son propiedades únicas de la sección."))
    for message in result.warnings:
        doc.add(Paragraph("Aviso: " + message))
    doc.add(Heading("Estados límite y eventos", 2))
    rows = []
    for record in result.limit_states:
        duct = result.ductility_of(record.curvature)
        rows.append([record.label, _num(None if not record.reached else record.curvature * PHI_UNIT, 2),
                     _num(None if not record.reached else record.moment / 1e6, 1), _num(duct, 1), record.definition])
    doc.add(Table(["Estado límite", "φ (µrad/mm)", "M (kN·m)", "μφ", "Criterio"], rows, widths=[2.2, 1.0, 1.0, 0.7, 4.0]))
    doc.add(Picture(moment_curvature_png(result), "Figura 2. Diagrama momento–curvatura y estados límite.", width_cm=15.5))
    doc.add(Picture(materials_png(model), "Figura 3. Curvas esfuerzo–deformación del concreto y del acero usadas en el cálculo.", width_cm=15.5))

    doc.add(Heading("3. Interpretación"))
    doc.add(Bullets([plain(line) for line in interpretation_lines(result)]))

    section_number = 4
    if interaction is not None and interaction.get("points"):
        from calculations.structural.column_interaction import interaction_dataframe

        doc.add(Heading(f"{section_number}. Diagrama de interacción P–M y ductilidad vs. carga axial"))
        doc.add(Picture(interaction_png(model, interaction["curve"], interaction["points"], interaction.get("current_axial_kn")),
                        "Figura 4. Izquierda: curva ideal (método de la norma) y momento máximo del análisis para cada carga axial. "
                        "Derecha: ductilidad de curvatura al terminar el análisis.", width_cm=16.0))
        curve = interaction["curve"]
        doc.add(Table(["Punto notable (método de la norma)", "P (kN)", "M (kN·m)"], [
            ["Tracción pura", _num(curve.tension_point[0], 1), _num(curve.tension_point[1], 1)],
            ["Punto balanceado", _num(curve.balanced_point[0], 1), _num(curve.balanced_point[1], 1)],
            ["Compresión pura", _num(curve.compression_point[0], 1), _num(curve.compression_point[1], 1)],
        ], widths=[3, 1.3, 1.3]))
        frame = interaction_dataframe(interaction["points"], model)
        doc.add(Table(["P (kN)", "P/f'co·Ag", "Mi (kN·m)", "Mmax (kN·m)", "φy (µrad/mm)", "μφ", "Termina por"], [
            [_num(r.axial_kN, 0), _num(r.axial_ratio, 3), _num(r.Mi_kNm, 1), _num(r.Mmax_kNm, 1),
             _num(r.phi_y_urad_per_mm, 2), _num(r.ductility, 1), r.termination or r.error]
            for r in frame.itertuples()], widths=[1, 1, 1, 1.1, 1.1, 0.8, 3]))
        doc.add(Paragraph("Donde la ductilidad aparece como «—» el acero no alcanzó la fluencia antes de terminar el análisis (respuesta controlada por compresión). "
                          "Cerca de la capacidad axial del modelo el equilibrio puede no converger; esos puntos se identifican en la columna «Termina por»."))
        section_number += 1

    if variants:
        from calculations.structural.column_variants import variants_dataframe

        doc.add(Heading(f"{section_number}. Comparación de configuraciones"))
        doc.add(Picture(variants_png(variants), "Figura 5. Curvas M–φ de las configuraciones comparadas.", width_cm=16.0))
        frame = variants_dataframe(variants)
        doc.add(Table(["Configuración", "s (mm)", "Ø est. (mm)", "P (kN)", "f'cc (MPa)", "Mmax (kN·m)", "φy (µrad/mm)", "μφ", "Termina por"], [
            [r.name, _num(getattr(r, "spacing_mm", None), 0), _num(getattr(r, "stirrup_diameter_mm", None), 1),
             _num(getattr(r, "axial_kN", None), 0), _num(getattr(r, "fcc_mpa", None), 1), _num(getattr(r, "Mmax_kNm", None), 1),
             _num(getattr(r, "phi_y_urad_per_mm", None), 2), _num(getattr(r, "ductility", None), 1), (getattr(r, "termination", "") if not r.error else r.error)]
            for r in frame.itertuples()], widths=[1.5, 0.8, 0.9, 0.9, 1, 1.1, 1.1, 0.7, 2.4]))
        section_number += 1

    doc.add(Heading(f"{section_number}. Alcance, hipótesis y referencias"))
    doc.add(Bullets([
        "Análisis por fibras con compatibilidad de deformaciones (secciones planas) y equilibrio axial; la curvatura se impone y la deformación axial se corrige por Newton (Mander, 1984, §7.4.2).",
        "Concreto del recubrimiento y del núcleo confinado con la curva de Popovics; confinamiento por el procedimiento de la superficie de resistencia (Mander, 1984, §3.4.5). Acero con meseta de fluencia y endurecimiento (Cap. 2).",
        "Respuesta monotónica sobre la curva esqueleto: no hay ramas cíclicas, descarga de fibras, efecto de velocidad de deformación, penetración de fluencia, cortante, deslizamiento ni pandeo de barras.",
        "Es una reconstrucción numérica de COLUMN basada en la literatura, no una traducción de su programa; se validó con las Tablas 5.1 y 8.3 de Soesianawati (1986).",
    ]))
    doc.add(Bullets([
        "Mander, J. B. (1984). Seismic Design of Bridge Piers. Research Report 84-2, Univ. of Canterbury.",
        "Mander, J. B., Priestley, M. J. N. y Park, R. (1988). Theoretical stress-strain model for confined concrete. J. Struct. Eng., 114(8), 1804-1826.",
        "Soesianawati, M. T. (1986). Limited Ductility Design of Reinforced Concrete Columns. Research Report 86-10, Univ. of Canterbury.",
    ]))
    return doc


# =========================================================================================
# Escritores
# =========================================================================================


def write_docx(doc: ReportDocument) -> bytes:
    """Vuelca el documento a Word (.docx) con python-docx."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt
    except ImportError as error:
        raise RuntimeError("Falta la librería python-docx: instálela con `pip install python-docx`.") from error

    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)

    document.add_heading(doc.title, 0)
    if doc.subtitle:
        subtitle = document.add_paragraph(doc.subtitle)
        subtitle.runs[0].italic = True
    for block in doc.blocks:
        if isinstance(block, Heading):
            document.add_heading(block.text, block.level)
        elif isinstance(block, Paragraph):
            document.add_paragraph(block.text)
        elif isinstance(block, Bullets):
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Light Grid Accent 1"
            for cell, header in zip(table.rows[0].cells, block.headers):
                cell.text = header
            for row in block.rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5)
            document.add_paragraph()
        elif isinstance(block, Picture):
            document.add_picture(BytesIO(block.image.data), width=Cm(block.width_cm))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if block.caption:
                caption = document.add_paragraph(block.caption)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].italic = True
                caption.runs[0].font.size = Pt(9)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _dejavu(name: str) -> str:
    """Ruta de una fuente DejaVu, que matplotlib distribuye y cubre griego, Ø, µ y ″."""
    import matplotlib

    return str(Path(matplotlib.get_data_path()) / "fonts" / "ttf" / name)


def write_pdf(doc: ReportDocument) -> bytes:
    """Vuelca el documento a PDF con fpdf2 (A4, fuente DejaVu para los símbolos)."""
    try:
        from fpdf import FPDF
    except ImportError as error:
        raise RuntimeError("Falta la librería fpdf2: instálela con `pip install fpdf2`.") from error

    class ReportPDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("DejaVu", "", 8)
            self.cell(0, 5, f"Página {self.page_no()}", align="C")

    pdf = ReportPDF(format="A4")
    pdf.set_auto_page_break(True, margin=16)
    pdf.set_margins(18, 16, 18)
    pdf.add_font("DejaVu", "", _dejavu("DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", _dejavu("DejaVuSans-Bold.ttf"))
    pdf.add_font("DejaVu", "I", _dejavu("DejaVuSans-Oblique.ttf"))
    pdf.add_page()
    usable = pdf.w - pdf.l_margin - pdf.r_margin

    def write(text, size, style="", height=5.4, align="L", indent=0.0):
        pdf.set_font("DejaVu", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(usable - indent, height, text, align=align, new_x="LMARGIN", new_y="NEXT")

    write(doc.title, 16, "B", 8, "C")
    if doc.subtitle:
        write(doc.subtitle, 9, "I", 5, "C")
    pdf.ln(3)
    for block in doc.blocks:
        if isinstance(block, Heading):
            pdf.ln(3 if block.level > 1 else 5)
            write(block.text, 13 if block.level == 1 else 11, "B", 6.5)
            pdf.ln(1)
        elif isinstance(block, Paragraph):
            write(block.text, 10)
            pdf.ln(1.5)
        elif isinstance(block, Bullets):
            for item in block.items:
                write("•  " + item, 10, indent=3)
            pdf.ln(1.5)
        elif isinstance(block, Table):
            widths = block.widths or [1] * len(block.headers)
            total = float(sum(widths))
            pdf.set_font("DejaVu", "", 8)
            with pdf.table(col_widths=tuple(w / total * usable for w in widths), line_height=4.4, text_align="LEFT",
                           width=usable) as table:
                header = table.row()
                for text in block.headers:
                    header.cell(str(text))
                for row in block.rows:
                    cells = table.row()
                    for value in row:
                        cells.cell(str(value))
            pdf.ln(3)
        elif isinstance(block, Picture):
            width = min(block.width_cm * 10.0, usable)
            height = width * block.image.aspect
            if pdf.get_y() + height + 12 > pdf.h - pdf.b_margin:
                pdf.add_page()
            pdf.image(BytesIO(block.image.data), x=pdf.l_margin + (usable - width) / 2.0, w=width)
            pdf.ln(1)
            if block.caption:
                write(block.caption, 8.5, "I", 4.5, "C")
            pdf.ln(2)
    return bytes(pdf.output())
